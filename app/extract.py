"""Extracción de texto de PDF, DOCX, PPTX, TXT/MD e imágenes (con OCR)."""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt", ".md", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}

# Si una página de PDF tiene menos caracteres que esto, se intenta OCR
MIN_CHARS_PER_PAGE = 40


class ExtractionError(Exception):
    """Error comprensible para el usuario."""


@dataclass
class Extraction:
    text: str
    pages: int = 0
    tables: list[list[list[str]]] = field(default_factory=list)
    used_ocr: bool = False
    warnings: list[str] = field(default_factory=list)

    @property
    def char_count(self) -> int:
        return len(self.text.strip())


def extract(path: Path) -> Extraction:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ExtractionError(
            f"El formato «{ext or 'sin extensión'}» todavía no está soportado. "
            f"Formatos aceptados: {', '.join(sorted(SUPPORTED))}."
        )
    try:
        if ext == ".pdf":
            result = _pdf(path)
        elif ext == ".docx":
            result = _docx(path)
        elif ext == ".pptx":
            result = _pptx(path)
        elif ext in {".txt", ".md"}:
            result = Extraction(text=path.read_text(encoding="utf-8", errors="replace"), pages=1)
        else:
            result = _image(path)
    except ExtractionError:
        raise
    except Exception as e:  # archivo dañado, protegido, etc.
        raise ExtractionError(f"No se ha podido leer el archivo ({type(e).__name__}: {e}).") from e

    if result.char_count < 200:
        result.warnings.append(
            "El documento tiene muy poco texto legible. Revisa que no esté escaneado con mala "
            "calidad: la app no inventará el contenido que falte."
        )
    return result


def _pdf(path: Path) -> Extraction:
    import pdfplumber

    parts, tables, used_ocr, warnings = [], [], False, []
    with pdfplumber.open(path) as pdf:
        n = len(pdf.pages)
        for i, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            for t in page.extract_tables() or []:
                tables.append([[c or "" for c in row] for row in t])
            if len(text.strip()) < MIN_CHARS_PER_PAGE:
                ocr = _ocr_pdf_page(page)
                if ocr is None:
                    warnings.append(f"Página {i}: sin texto y el OCR no está disponible.")
                elif len(ocr.strip()) >= MIN_CHARS_PER_PAGE:
                    text, used_ocr = ocr, True
                else:
                    warnings.append(f"Página {i}: no se ha podido leer el texto (ni con OCR).")
            parts.append(text)
    return Extraction(text="\n\n".join(parts), pages=n, tables=tables, used_ocr=used_ocr, warnings=warnings)


def _ocr_pdf_page(page) -> str | None:
    try:
        import pytesseract
        img = page.to_image(resolution=300).original
        return pytesseract.image_to_string(img, lang=_ocr_langs())
    except Exception:
        return None


def _ocr_langs() -> str:
    """Usa español y catalán si están instalados; si no, inglés."""
    try:
        import pytesseract
        available = set(pytesseract.get_languages(config=""))
    except Exception:
        return "eng"
    wanted = [l for l in ("spa", "cat", "eng") if l in available]
    return "+".join(wanted) or "eng"


def _docx(path: Path) -> Extraction:
    import docx

    d = docx.Document(str(path))
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style.startswith("título"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            lines.append("#" * min(int(level) + 1, 4) + " " + t)
        else:
            lines.append(t)
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    return Extraction(text="\n\n".join(lines), pages=1, tables=tables)


def _pptx(path: Path) -> Extraction:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        texts.append(t)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        block = f"## Diapositiva {i}\n" + "\n".join(texts)
        if notes:
            block += f"\n\nNotas: {notes}"
        parts.append(block)
    return Extraction(text="\n\n".join(parts), pages=len(prs.slides))


def _image(path: Path) -> Extraction:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as e:
        raise ExtractionError("El OCR no está disponible en este equipo.") from e
    img = Image.open(io.BytesIO(path.read_bytes()))
    text = pytesseract.image_to_string(img, lang=_ocr_langs())
    return Extraction(text=text, pages=1, used_ocr=True)
